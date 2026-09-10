from openpyxl import load_workbook


CLUSTER_COL_MAP = {
    "HL3": 0,
    "MOBILE-CONTROL": 1,
    "MOBILE-CONTROL-RD": 2,
    "MOBILE-CONTROL-HRT-export": 3,
    "MOBILE-CONTROL-SRT-import": 4,
    "MOBILE-DATA": 5,
    "MOBILE-DATA-RD": 6,
    "MOBILE-DATA-HRT-export": 7,
    "MOBILE-DATA-SRT-import": 8,
    "MOBILE-ACCESS-MGMT": 9,
    "MOBILE-ACCESS-MGMT-RD": 10,
    "MOBILE-ACCESS-MGMT-SRT-import": 11,
    "MOBILE-ACCESS-MGMT-HRT-export": 12,
    "DHCP 1": 13,
    "DHCP 2": 14,
    "DHCP 3": 15,
    "UF": 16,
    "RF VENDOR": 17,
}


def _get_value(row, index):
    if index < len(row):
        return row[index]
    return None


def read_cluster_excel(file_path: str):
    wb = load_workbook(file_path, data_only=True)
    ws = wb["CLUSTER"]

    data = []

    # linha 3 é exemplo/dado inicial
    for row in ws.iter_rows(min_row=3, values_only=True):
        row_dict = {field: _get_value(row, idx) for field, idx in CLUSTER_COL_MAP.items()}
        data.append(row_dict)

    return data