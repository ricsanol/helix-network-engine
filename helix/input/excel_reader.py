from openpyxl import load_workbook


REQUIRED_COLUMNS = [
    "Hostname",
    "Vendor",
    "Modelo",
    "Loopback /32 IPV4",
    "RR1",
    "ID_SIGLA",
    "INTERFACE-SERV",   # ← NOVO
    "GATEWAY MOBILE-DATA",
    "BASEBAND MOBILE-DATA",
    "VLAN MOBILE-DATA",
    "GATEWAY MOBILE-CONTROL",
    "BASEBAND MOBILE-CONTROL",
    "VLAN MOBILE-CONTROL",
    "GATEWAY MOBILE-MGMT",
    "BASEBAND MOBILE-MGMT",
    "VLAN MOBILE-MGMT",
]


def _normalize_header(value):
    if value is None:
        return ""
    return str(value).strip()


def read_excel(file_path: str):
    wb = load_workbook(file_path, data_only=True)
    ws = wb.active

    # Linha 3 = cabeçalhos
    header_cells = list(ws.iter_rows(min_row=3, max_row=3, values_only=True))[0]
    headers = [_normalize_header(cell) for cell in header_cells]

    header_index = {header: idx for idx, header in enumerate(headers) if header}

    missing = [col for col in REQUIRED_COLUMNS if col not in header_index]
    if missing:
        raise ValueError(f"Colunas obrigatórias não encontradas na planilha: {missing}")

    data = []

    # Dados começam na linha 4
    for row in ws.iter_rows(min_row=4, values_only=True):
        row_dict = {}
        for col in REQUIRED_COLUMNS:
            idx = header_index[col]
            row_dict[col] = row[idx] if idx < len(row) else None
        data.append(row_dict)

    return data